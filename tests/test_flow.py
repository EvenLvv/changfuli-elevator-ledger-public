import io
import os
import tempfile
import unittest
import zipfile
from datetime import date, timedelta
from pathlib import Path

from PIL import Image


class FlowTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from support import isolated_app
        cls.context = isolated_app()
        cls.client, cls.paths = cls.context.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls.context.__exit__(None, None, None)

    def test_complete_flow(self):
        c = self.client
        def post(path, payload):
            response = c.post(path, json=payload)
            self.assertEqual(response.status_code, 200, response.text)
            return response.json()

        self.assertTrue(c.get('/api/setup/status').json()['needed'])
        from fastapi.testclient import TestClient
        from backend.main import app
        remote = TestClient(app, client=('192.168.1.10', 50000))
        self.assertEqual(remote.post('/api/setup', json={'username':'intruder','display_name':'外部','password':'a-long-pass-123'}).status_code, 403)
        post('/api/setup', {'username':'admin','display_name':'测试管理员','password':'a-long-pass-123'})
        post('/api/login', {'username':'admin','password':'a-long-pass-123'})
        building = post('/api/buildings', {'name':'18栋'})['id']
        unit = post('/api/units', {'building_id':building,'name':'1单元'})['id']
        elevator = post('/api/elevators', {'unit_id':unit,'name':'1号梯','equipment_number':'T-001',
            'inspection_on':(date.today()+timedelta(days=6)).isoformat(),
            'maintenance_due_on':(date.today()-timedelta(days=1)).isoformat()})['id']
        for typ in ('日常维保','故障','维修'):
            post('/api/events', {'elevator_id':elevator,'type':typ,'title':typ+'门锁检查',
                                 'occurred_on':date.today().isoformat(),'work_content':'检查门锁', 'cost_yuan':'10.50',
                                 'elevator_status':'故障' if typ=='故障' else '正常' if typ=='维修' else ''})
            if typ == '故障':
                self.assertEqual(c.get(f'/api/elevators/{elevator}').json()['status'],'故障')
        past = (date.today()-timedelta(days=370)).isoformat()
        first = post('/api/events', {'elevator_id':elevator,'type':'配件更换','title':'更换门锁开关',
            'occurred_on':past,'parts':[{'name':'门锁开关','quantity':1,'unit_price_yuan':'260',
                                         'warranty_days':365,'replaced_on':past}]})['id']
        current = date.today().isoformat()
        second = post('/api/events', {'elevator_id':elevator,'type':'配件更换','title':'再次更换门锁开关',
            'occurred_on':current,'parts':[{'name':'门锁开关','quantity':1,'unit_price_yuan':'270',
                                            'warranty_days':7,'replaced_on':current}]})['id']
        self.assertEqual(c.get(f'/api/events/{first}').json()['cost_cents'],26000)
        self.assertEqual(c.get(f'/api/events/{second}').json()['parts'][0]['total_price_cents'],27000)
        image = Image.new('RGB',(1200,600),'blue'); image_bytes=io.BytesIO(); image.save(image_bytes,format='JPEG')
        file_ids = []
        for filename, content, kind in (('现场.jpg',image_bytes.getvalue(),'image/jpeg'),
                                        ('现场.jpg',image_bytes.getvalue(),'image/jpeg'),
                                        ('报价.pdf',b'%PDF-1.4\n%%EOF','application/pdf'),
                                        ('签字.docx',self.office_docx(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
                                        ('清单.xlsx',self.office_xlsx(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')):
            response = c.post(f'/api/events/{second}/attachments',files={'file':(filename,content,kind)},
                              data={'category':'工作现场照片' if filename.endswith('jpg') else '其他'})
            self.assertEqual(response.status_code,200,response.text)
            file_ids.append(response.json()['id'])
            stored = c.get('/api/attachments/'+response.json()['id'])
            self.assertEqual(stored.content,content)
        self.assertEqual(len(file_ids),len(set(file_ids)))
        self.assertEqual(len(c.get(f'/api/events?elevator_id={elevator}').json()),5)
        first_page=c.get('/api/events',params={'elevator_id':elevator,'limit':2}).json()
        second_page=c.get('/api/events',params={'elevator_id':elevator,'limit':2,'offset':2}).json()
        self.assertEqual(len(first_page),2)
        self.assertFalse({x['id'] for x in first_page} & {x['id'] for x in second_page})
        parts = c.get(f'/api/elevators/{elevator}/parts').json()
        self.assertEqual(parts[0]['days_since_previous'],370)
        self.assertEqual(parts[0]['warranty_until'],(date.today()+timedelta(days=7)).isoformat())
        for q in ('18栋','1号梯','门锁','维修'):
            self.assertTrue(c.get('/api/search',params={'q':q}).json()['elevators'],q)
        self.assertTrue(c.get('/api/search',params={'q':'18栋 门锁'}).json()['events'])
        kinds = [x['kind'] for x in c.get('/api/reminders').json()]
        self.assertIn('配件质保',kinds);self.assertIn('维保',kinds);self.assertIn('年检',kinds)
        self.assertEqual(c.get('/api/dashboard').json()['stats']['total'],1)
        self.assertEqual(c.get(f'/api/elevators/{elevator}').json()['status'],'正常')
        self.assertEqual(c.get('/api/costs',params={'building_id':building}).json()['cost_cents'],56150)
        self.assertEqual(c.put('/api/network',json={'base_url':'http://192.168.1.20:8765'}).status_code,200)
        self.assertEqual(c.get('/api/elevators/'+str(elevator)+'/qr').headers['content-type'],'image/png')
        for scope, params in (('elevator',{'elevator_id':elevator}),('building',{'building_id':building})):
            response=c.get('/api/export',params={'scope':scope,**params,'start':past,'end':current})
            self.assertEqual(response.status_code,200,response.text[:200] if response.status_code!=200 else '')
            from docx import Document
            document=Document(io.BytesIO(response.content))
            text='\n'.join(p.text for p in document.paragraphs)
            self.assertIn('门锁开关',text)
            self.assertIn('电梯工作记录档案',text)
            self.assertEqual(len(document.inline_shapes),2)
            self.assertAlmostEqual(document.inline_shapes[0].width/document.inline_shapes[0].height,2.0,delta=0.1)
        empty=c.get('/api/export',params={'scope':'elevator','elevator_id':elevator,'start':'2000-01-01','end':'2000-12-31'})
        self.assertEqual(empty.status_code,200)
        self.assertIn('暂无工作记录', '\n'.join(p.text for p in Document(io.BytesIO(empty.content)).paragraphs))
        package=c.get('/api/export',params={'scope':'elevator','elevator_id':elevator,'package':'true'})
        self.assertEqual(package.status_code,200)
        with zipfile.ZipFile(io.BytesIO(package.content)) as archive:
            self.assertTrue(any(n.endswith('.docx') for n in archive.namelist()))
            self.assertTrue(any(n.endswith('.jpg') for n in archive.namelist()))
            self.assertTrue(any(n.endswith('.pdf') for n in archive.namelist()))
        prepared=post('/api/export/prepare', {'scope':'elevator','elevator_id':elevator,
                                               'start':past,'end':current,'package':True})
        self.assertTrue(prepared['filename'].endswith('.zip'))
        self.assertEqual(c.get(prepared['download']).status_code,200)
        self.assertEqual(c.post('/api/export/prepare',json={'scope':'elevator'}).status_code,400)
        backup=post('/api/backup',{})
        archive=c.get(backup['download'])
        self.assertEqual(archive.status_code,200)
        with zipfile.ZipFile(io.BytesIO(archive.content)) as content:
            self.assertIn('data/elevator.db',content.namelist())
            self.assertTrue(any(n.startswith('storage/photos/') for n in content.namelist()))
        self.assertTrue(c.get('/api/audit').json())
        self.assertEqual(c.post('/api/events',json={'elevator_id':elevator,'type':'维修','title':'坏金额',
          'occurred_on':current,'cost_yuan':'12x'}).status_code,400)
        self.assertEqual(c.post(f'/api/events/{first}/attachments',files={'file':('../bad.exe',b'bad','application/octet-stream')}).status_code,400)
        free_part = post('/api/events', {'elevator_id':elevator,'type':'配件更换','title':'质保内免费更换',
            'occurred_on':current,'cost_yuan':'0','parts':[{'name':'免费配件','quantity':1,
                'unit_price_yuan':'100','replaced_on':current}]})['id']
        self.assertEqual(c.get(f'/api/events/{free_part}').json()['cost_cents'],0)
        staff = post('/api/users', {'username':'staff','display_name':'工作人员','password':'staff-pass-12345','role':'staff'})['id']
        self.assertEqual(c.put(f'/api/users/{staff}',json={'display_name':'新工作人员','active':True}).status_code,200)
        self.assertEqual(c.put('/api/users/1',json={'role':'staff'}).status_code,400)
        from fastapi.testclient import TestClient
        from backend.main import app
        other = TestClient(app)
        self.assertEqual(other.post('/api/login',json={'username':'staff','password':'staff-pass-12345'}).status_code,200)
        self.assertEqual(other.post('/api/buildings',json={'name':'无权限楼栋'}).status_code,403)
        self.assertEqual(other.put(f'/api/events/{first}',json={'type':'维修','title':'越权','occurred_on':current}).status_code,403)
        self.assertEqual(c.put(f'/api/buildings/{building}',json={'name':'18号楼'}).status_code,200)
        self.assertEqual(c.put(f'/api/units/{unit}',json={'name':'第一单元'}).status_code,200)

    @staticmethod
    def office_docx():
        from docx import Document
        output=io.BytesIO();Document().save(output);return output.getvalue()

    @staticmethod
    def office_xlsx():
        from openpyxl import Workbook
        output=io.BytesIO();Workbook().save(output);return output.getvalue()


if __name__ == '__main__':
    unittest.main()
