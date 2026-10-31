"""Editable archives: isolated source snapshots, ordering, permissions, and recovery."""
import csv
import io
import json
import os
import subprocess
import sys
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
from support import isolated_app


class ExportDraftTest(unittest.TestCase):
    def setUp(self):
        self.context = isolated_app()
        self.client, self.paths = self.context.__enter__()
        self.addCleanup(self.context.__exit__, None, None, None)
        self.post('/api/setup', {'username': 'admin', 'display_name': '管理员', 'password': 'draft-test-12345'})
        self.post('/api/login', {'username': 'admin', 'password': 'draft-test-12345'})
        self.lift = self.post('/api/elevators', {'building_name': '18栋', 'unit_name': '1单元', 'name': '1号梯',
                              'equipment_number': '00018', 'maintenance_company': '物业维保单位', 'maintenance_phone': '00123'})['id']

    def post(self, url, data, client=None):
        response = (client or self.client).post(url, json=data)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def event(self, **values):
        return self.post('/api/events', {'elevator_id': self.lift, 'title': '门锁维修', 'type': '维修',
                                         'occurred_on': '2026-10-28', **values})['id']

    def photo(self, event, color, name='现场.png', fmt='PNG'):
        raw = io.BytesIO(); Image.new('RGB', (120, 80), color).save(raw, format=fmt)
        response = self.client.post(f'/api/events/{event}/attachments', files={'file': (name, raw.getvalue())},
                                    data={'category': '工作现场照片'})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()['id'], raw.getvalue()

    def draft(self, **filters):
        return self.post('/api/export-drafts', {'scope': 'elevator', 'elevator_id': self.lift, **filters})

    def save(self, draft, **values):
        response = self.client.put(f"/api/export-drafts/{draft['id']}", json={**draft, **values})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def prepare(self, draft, **options):
        return self.post(f"/api/export-drafts/{draft['id']}/prepare", {'revision': draft['revision'], **options})

    def test_default_natural_paragraphs_metadata_filters_and_all_source_fields(self):
        empty = self.draft(directory=False, package=True)
        self.assertLessEqual(len(empty['blocks']), 6)
        self.assertFalse(any('—' in b.get('text', '') for b in empty['blocks']))
        self.assertEqual(empty['metadata'], {'community': '昌福里社区27区党群物业服务站',
                                           'scope_label': '18栋 1单元 1号梯', 'period_label': '全部历史资料'})
        self.assertFalse(empty['filters']['directory']); self.assertTrue(empty['filters']['package'])
        event = self.event(type='故障', completed=False, stopped=True, recovered=False,
                           fault_at='2026-10-28T08:00', reported_at='2026-10-28T08:02', arrived_at='2026-10-28T08:20',
                           started_at='2026-10-28T08:21', ended_at='2026-10-28T09:00', restored_at='2026-10-28T09:01',
                           work_content='检查门锁\n测量接点', symptom='无法关门', cause='接点磨损', remedy='换门锁',
                           resolution='现场试运行', workers='张师傅', company='维修单位甲', notes='复查记录', cost_yuan='0',
                           parts=[{'name': '门锁', 'model': 'MK-1', 'quantity': 1, 'unit_price_yuan': '0',
                                   'brand': '配件厂家', 'reason': '接点磨损', 'notes': '保留旧件', 'warranty_days': 0}])
        draft = self.draft(start='2026-10-01', end='2026-10-28')
        text = '\n'.join(b.get('text', '') for b in draft['blocks'])
        for value in ['00018', '00123', '未完成', '发生停梯：是', '尚未恢复', '08:00', '08:02', '08:20', '08:21', '09:00',
                      '09:01', '检查门锁\n测量接点', '无法关门', '接点磨损', '换门锁', '现场试运行', '张师傅', '维修单位甲',
                      '复查记录', 'MK-1', '配件厂家', '质保天数：0', '质保开始日期：2026-10-28', '0.00元', '保留旧件']:
            self.assertIn(value, text)
        self.assertLess(len(draft['blocks']), 22)
        self.assertIn('2026-10-01', draft['title'])
        self.assertIn('18栋', draft['title'])
        document = Document(io.BytesIO(self.client.get(self.prepare(draft)['download']).content))
        self.assertIn('编制单位：昌福里社区27区党群物业服务站', '\n'.join(p.text for p in document.paragraphs))
        self.assertIn('统计周期：2026-10-01 至 2026-10-28', [p.text for p in document.paragraphs])
        self.assertEqual(next(p for p in document.paragraphs if p.text.startswith('电梯：')).alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(self.client.get(f'/api/events/{event}').json()['title'], '门锁维修')

    def test_edited_order_docx_zip_catalog_and_original_records_unchanged(self):
        from backend.db import connect
        event = self.event(workers='维修人员')
        first, first_raw = self.photo(event, 'red', 'cgi-bin_mmwebwx-bin_webwxgetmsgimg__&MsgID=777&skey=@crypt_long.png')
        second, second_raw = self.photo(event, 'green', '绿色.webp', 'WEBP')
        removed, removed_raw = self.photo(event, 'blue')
        self.assertEqual(self.client.post(f'/api/events/{event}/attachments', files={'file': ('签字资料.pdf', b'%PDF-1.4\noriginal')}).status_code, 200)
        before = self.client.get(f'/api/events/{event}').json()
        draft = self.draft()
        with connect() as db:
            snapshot = db.execute('SELECT source_json FROM export_drafts WHERE id=?', (draft['id'],)).fetchone()[0]
        blocks = [{'id': 'a', 'type': 'paragraph', 'text': '编辑后的开场正文\r补充说明'},
                  {'id': 'b', 'type': 'image', 'attachment_id': second, 'caption': '调整后第一张'},
                  {'id': 'c', 'type': 'heading', 'text': '自定小标题'},
                  {'id': 'd', 'type': 'image', 'attachment_id': first, 'caption': '调整后第二张'},
                  {'id': 'e', 'type': 'paragraph', 'text': '编辑后的结束正文'}]
        draft = self.save(draft, title='更名后的检查文稿', blocks=blocks, source_json='malicious', metadata={'community': '伪造'})
        self.assertEqual(draft['blocks'], blocks)
        self.assertEqual([e['id'] for e in draft['attachment_directory'][:2]], [second, first])
        self.assertEqual([e['reference'] for e in draft['attachment_directory'][:3]], ['图1', '图2', '未编入正文'])
        self.assertTrue(all('cgi-bin' not in e['display_name'] for e in draft['attachment_directory']))
        output = self.prepare(draft, package=True)
        with zipfile.ZipFile(io.BytesIO(self.client.get(output['download']).content)) as archive:
            self.assertIsNone(archive.testzip())
            rows = list(csv.DictReader(io.StringIO(archive.read('原始附件清单.csv').decode('utf-8-sig'))))
            self.assertEqual([row['资料编号'] for row in rows], ['图1', '图2', '未编入正文', '附件1'])
            for row, raw in zip(rows, [second_raw, first_raw, removed_raw, b'%PDF-1.4\noriginal']):
                self.assertEqual(archive.read(row['包内路径']), raw)
            doc = Document(io.BytesIO(archive.read(next(n for n in archive.namelist() if n.endswith('.docx')))))
            text = '\n'.join(p.text for p in doc.paragraphs)
            self.assertEqual(next(p for p in doc.paragraphs if p.text.startswith('编辑后的开场正文')).alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assertLess(text.index('编辑后的开场正文'), text.index('图1  调整后第一张'))
            self.assertLess(text.index('图1  调整后第一张'), text.index('自定小标题'))
            self.assertLess(text.index('自定小标题'), text.index('图2  调整后第二张'))
            self.assertLess(text.index('图2  调整后第二张'), text.index('编辑后的结束正文'))
            self.assertEqual(len(doc.inline_shapes), 2)
            cells = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            self.assertNotIn('cgi-bin', text + cells); self.assertIn('未编入正文', cells)
            self.assertEqual(doc.tables[-1].rows[1].cells[1].text, '18栋 1单元 1号梯\n门锁维修')
        self.assertEqual(self.client.get(f'/api/events/{event}').json(), before)
        with connect() as db:
            self.assertEqual(db.execute('SELECT source_json FROM export_drafts WHERE id=?', (draft['id'],)).fetchone()[0], snapshot)
        self.assertEqual(self.client.get('/api/attachments/' + removed).content, removed_raw)

    def test_validation_snapshot_whitelist_and_revision_are_required(self):
        event = self.event(); photo, _ = self.photo(event, 'red')
        excluded = self.event(occurred_on='2026-09-01'); other, _ = self.photo(excluded, 'blue')
        draft = self.draft(start='2026-10-01')
        endpoint = f"/api/export-drafts/{draft['id']}"
        image = {'id': 'image', 'type': 'image', 'attachment_id': photo, 'caption': ''}
        invalid = [[{**image, 'attachment_id': other}], [image, {**image, 'id': 'second'}],
                   [image, image], [{'id': '../bad', 'type': 'paragraph', 'text': 'x'}],
                   [{'id': 'bad', 'type': 'html', 'text': 'x'}], [{'id': 'bad', 'type': 'paragraph', 'text': '\x00'}],
                   [{'id': 'bad', 'type': 'heading', 'text': 'x' * 301}], [None]]
        for blocks in invalid:
            response = self.client.put(endpoint, json={'title': '校验文稿', 'revision': 1, 'blocks': blocks})
            self.assertEqual(response.status_code, 400, response.text)
        for revision in [None, True, '1.0', 0]:
            self.assertEqual(self.client.put(endpoint, json={'title': '校验文稿', 'revision': revision, 'blocks': []}).status_code, 400)
        for changes in [{'package': 'false'}, {'directory': []}, {'photos': 'yes'}, {'scope': []}, {'start': 'bad'}]:
            self.assertEqual(self.client.post('/api/export-drafts', json=changes).status_code, 400)
        self.assertEqual(self.client.get(endpoint).json()['revision'], 1)

    def test_concurrent_save_has_exactly_one_winner_and_stale_export_fails(self):
        draft = self.draft(); endpoint = f"/api/export-drafts/{draft['id']}"
        def save(title):
            return self.client.put(endpoint, json={**draft, 'title': title}).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(sorted(pool.map(save, ['甲稿', '乙稿'])), [200, 409])
        current = self.client.get(endpoint).json()
        self.assertEqual(current['revision'], 2)
        self.assertEqual(self.client.post(endpoint + '/prepare', json={'revision': 1}).status_code, 409)
        self.assertEqual(len([e for e in self.client.get('/api/audit').json() if e['target'] == '工作文稿' and e['action'] == '修改']), 1)

    def test_owner_admin_permissions_and_no_legacy_download_bypass(self):
        from backend.main import app
        from fastapi.testclient import TestClient
        event = self.event(); photo, _ = self.photo(event, 'red'); draft = self.draft()
        first = self.prepare(draft); second = self.prepare(draft)
        for output in [first, second]:
            stored = output['download'].rsplit('/', 1)[1]
            self.assertEqual(self.client.get('/api/exports/' + stored).status_code, 400)
        self.assertEqual(self.client.get(first['download']).status_code, 404)
        self.post('/api/users', {'username': 'staff', 'display_name': '工作人员', 'password': 'draft-staff-12345', 'role': 'staff'})
        with TestClient(app) as other:
            self.assertEqual(other.get(second['download']).status_code, 401)
            self.post('/api/login', {'username': 'staff', 'password': 'draft-staff-12345'}, other)
            self.assertEqual(other.get('/api/export-drafts').json(), [])
            endpoint = f"/api/export-drafts/{draft['id']}"
            for response in [other.get(endpoint), other.put(endpoint, json=draft),
                             other.post(endpoint + '/prepare', json={'revision': 1}), other.get(second['download']),
                             other.get(f'{endpoint}/images/{photo}')]:
                self.assertEqual(response.status_code, 403)
            own = self.post('/api/export-drafts', {'scope': 'all'}, other)
            self.assertEqual(len(other.get('/api/export-drafts').json()), 1)
        self.assertEqual(len(self.client.get('/api/export-drafts').json()), 2)
        self.save(own, title='管理员可代编辑')
        draft = self.save(draft, title='正文更新后原下载失效')
        self.assertEqual(self.client.get(second['download']).status_code, 404)

    def test_every_parent_soft_delete_and_missing_file_block_exports(self):
        from backend.db import connect
        event = self.event(); photo, raw = self.photo(event, 'red'); draft = self.draft()
        endpoint = f"/api/export-drafts/{draft['id']}"
        with connect() as db:
            lift = dict(db.execute('SELECT * FROM elevators WHERE id=?', (self.lift,)).fetchone())
            unit = dict(db.execute('SELECT * FROM units WHERE id=?', (lift['unit_id'],)).fetchone())
            relative = db.execute('SELECT relative_path FROM attachments WHERE id=?', (photo,)).fetchone()[0]
        for table, record_id in [('attachments', photo), ('events', event), ('elevators', self.lift),
                                 ('units', unit['id']), ('buildings', unit['building_id'])]:
            with self.subTest(table=table):
                with connect() as db:
                    db.execute(f'UPDATE {table} SET deleted_at=? WHERE id=?', ('2026-10-28', record_id))
                detail = self.client.get(endpoint).json()
                self.assertTrue(detail['source_warnings']); self.assertFalse(detail['images'][0]['available'])
                self.assertEqual(self.client.get(detail['images'][0]['url']).status_code, 404)
                self.assertEqual(self.client.post(endpoint + '/prepare', json={'revision': 1}).status_code, 400)
                with connect() as db:
                    db.execute(f'UPDATE {table} SET deleted_at=NULL WHERE id=?', (record_id,))
        path = self.paths['STORAGE'] / relative; path.unlink()
        self.assertEqual(self.client.post(endpoint + '/prepare', json={'revision': 1}).status_code, 400)
        self.assertFalse(list(self.paths['EXPORTS'].iterdir()))
        path.write_bytes(raw)
        self.assertTrue(self.client.get(endpoint).json()['images'][0]['available'])

    def test_empty_elevator_deleted_and_single_image_validation_is_bounded(self):
        from backend.db import connect
        from backend import main
        empty = self.draft()
        with connect() as db:
            db.execute('UPDATE elevators SET deleted_at=? WHERE id=?', ('2026-10-28', self.lift))
        self.assertTrue(self.client.get(f"/api/export-drafts/{empty['id']}").json()['source_warnings'])
        self.assertEqual(self.client.post(f"/api/export-drafts/{empty['id']}/prepare", json={'revision': 1}).status_code, 400)
        with connect() as db:
            db.execute('UPDATE elevators SET deleted_at=NULL WHERE id=?', (self.lift,))
        for _ in range(3):
            event = self.event(); self.photo(event, 'red'); self.photo(event, 'green')
        draft = self.draft()
        with patch('backend.main.source_status', wraps=main.source_status) as check:
            self.assertEqual(self.client.get(draft['images'][0]['url']).status_code, 200)
            checked = check.call_args[0][0]
            self.assertEqual(len(checked), 1); self.assertEqual(len(checked[0]['events']), 1)
            self.assertEqual(len(checked[0]['events'][0]['attachments']), 1)

    def test_backup_restores_draft_snapshot_images_and_can_export_again(self):
        event = self.event(); self.photo(event, 'red')
        draft = self.save(self.draft(), title='恢复后文稿')
        self.prepare(draft)
        output = self.post('/api/backup', {})
        restored = self.paths['DATA'].parent / 'restored'
        with zipfile.ZipFile(io.BytesIO(self.client.get(output['download']).content)) as archive:
            archive.extractall(restored)
        environment = os.environ.copy()
        for key, folder in [('DATA', 'data'), ('STORAGE', 'storage'), ('EXPORT', 'exports'), ('BACKUP', 'backups'), ('LOG', 'logs')]:
            environment[f'ELEVATOR_{key}_DIR'] = str(restored / folder)
        code = '''from backend.main import app
from fastapi.testclient import TestClient
with TestClient(app) as c:
 assert c.get('/api/export-drafts').status_code == 401
 assert c.post('/api/login',json={'username':'admin','password':'draft-test-12345'}).status_code == 200
 draft=c.get('/api/export-drafts').json()[0]
 assert draft['title']=='恢复后文稿' and draft['revision']==2
 detail=c.get('/api/export-drafts/'+str(draft['id'])).json()
 assert detail['images'][0]['available']
 output=c.post('/api/export-drafts/'+str(draft['id'])+'/prepare',json={'revision':2,'package':True})
 assert output.status_code==200,output.text
 assert c.get(output.json()['download']).status_code==200
 print('restored draft snapshot, source image and export passed')
'''
        result = subprocess.run([sys.executable, '-c', code], env=environment, capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


if __name__ == '__main__':
    unittest.main()
