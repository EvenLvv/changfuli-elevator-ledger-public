"""Regression for inspection archives and saved attachment order; isolated data only."""
import csv
import io
import unittest
import zipfile
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from support import isolated_app


class ExportFormatTest(unittest.TestCase):
    def setUp(self):
        self.context = isolated_app()
        self.client, self.paths = self.context.__enter__()
        self.addCleanup(self.context.__exit__, None, None, None)
        self.post('/api/setup', {'username':'admin','display_name':'管理员','password':'format-password-123'})
        self.post('/api/login', {'username':'admin','password':'format-password-123'})
        self.lift = self.post('/api/elevators', {'building_name':'18栋','unit_name':'1单元','name':'1号梯'})['id']

    def post(self, url, data):
        r = self.client.post(url, json=data)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def event(self, title, time):
        return self.post('/api/events', {'elevator_id':self.lift,'type':'维修','title':title,
                         'occurred_on':'2026-09-30','started_at':time})['id']

    def photo(self, event, name, color):
        raw = io.BytesIO(); Image.new('RGB', (80, 140), color).save(raw, format='PNG')
        r = self.client.post(f'/api/events/{event}/attachments', files={'file':(name, raw.getvalue(), 'image/png')},
                             data={'category':'工作现场照片'})
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()['id'], raw.getvalue()

    def test_numbering_chronology_originals_and_formal_layout(self):
        from backend.exporter import safe_name
        long_path = self.paths['EXPORTS'] / ('0123456789ab_' + safe_name('电梯档案' * 100, 180) + '.docx')
        long_path.write_bytes(b'filename boundary check')
        self.assertLessEqual(len(long_path.name.encode('utf-8')), 255)
        self.assertTrue(long_path.is_file())
        # Late work is entered first, to expose insertion-order regressions.
        late = self.event('午后维修', '2026-09-30T14:00')
        early = self.event('上午检查', '2026-09-30T09:00')
        names = ['cgi-bin_mmwebwx-bin_webwxgetmsgimg__&MsgID=777&skey=@crypt_abc.png', '照片二.png']
        first, first_raw = self.photo(early, names[0], 'red')
        second, second_raw = self.photo(early, names[1], 'green')
        third, third_raw = self.photo(late, '照片三.png', 'blue')
        self.assertEqual(self.client.put(f'/api/events/{early}/attachment-order', json={'ids':[second,first]}).status_code,200)
        empty_lift = self.post('/api/elevators', {'building_name':'18栋','unit_name':'1单元','name':'2号梯'})['id']
        response = self.client.get('/api/export', params={'scope':'all'})
        self.assertEqual(response.status_code, 200, response.text if response.status_code!=200 else '')
        doc = Document(io.BytesIO(response.content))
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertLess(text.index('工作事项：上午检查'), text.index('工作事项：午后维修'))
        self.assertIn('暂无工作记录', text)
        cells = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for reference in ['图1-1','图1-2','图2-1']:
            self.assertIn(reference, cells)
        self.assertNotIn('cgi-bin', cells)
        self.assertNotIn('照片二.png', cells)
        self.assertIn('昌福里社区27区党群物业服务站', text + cells)
        for section in doc.sections:
            self.assertAlmostEqual(section.top_margin.cm, 3.7, places=2)
            self.assertAlmostEqual(section.bottom_margin.cm, 3.5, places=2)
            self.assertAlmostEqual(section.left_margin.cm, 2.8, places=2)
            self.assertAlmostEqual(section.right_margin.cm, 2.6, places=2)
        self.assertEqual(doc.styles['Normal'].font.size.pt, 16)
        self.assertEqual(doc.styles['Normal']._element.rPr.rFonts.get(qn('w:eastAsia')), '仿宋')
        self.assertEqual(doc.styles['Title']._element.rPr.rFonts.get(qn('w:eastAsia')), '方正小标宋简体')
        self.assertEqual(len(doc.inline_shapes), 3)
        self.assertLessEqual(doc.inline_shapes[0].width.cm, 7.01)
        self.assertLessEqual(doc.inline_shapes[0].height.cm, 6.21)
        rel = doc.part.related_parts[doc.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed]
        with Image.open(io.BytesIO(rel.blob)) as embedded:
            self.assertEqual(embedded.getpixel((0,0)), (0,128,0))
        self.assertEqual(len(doc.sections), 2)
        self.assertFalse(doc.sections[0].footer.paragraphs[0].text)
        self.assertIn('PAGE', doc.sections[1].footer._element.xml)
        package = self.client.get('/api/export', params={'scope':'all','package':True})
        with zipfile.ZipFile(io.BytesIO(package.content)) as archive:
            manifest = list(csv.DictReader(io.StringIO(archive.read('原始附件清单.csv').decode('utf-8-sig'))))
            self.assertEqual([r['资料编号'] for r in manifest], ['图1-1','图1-2','图2-1'])
            self.assertEqual(archive.read(manifest[0]['包内路径']), second_raw)
            self.assertEqual(archive.read(manifest[1]['包内路径']), first_raw)
            self.assertEqual(manifest[1]['原始文件名'], names[0])
            self.assertEqual(archive.read(manifest[2]['包内路径']), third_raw)
            self.assertFalse(any('cgi-bin' in name for name in archive.namelist()))
        filtered = self.client.get('/api/export', params={'scope':'elevator','elevator_id':self.lift,'photos':False,'directory':False})
        self.assertEqual(len(Document(io.BytesIO(filtered.content)).inline_shapes),0)

    def test_order_is_authorized_complete_and_migrates_old_schema(self):
        event = self.event('处理顺序', '2026-09-30T09:00')
        first, _ = self.photo(event, '一.png', 'red'); second, _ = self.photo(event, '二.png', 'blue')
        endpoint = f'/api/events/{event}/attachment-order'
        for invalid in [None, [first], [first, first], [first, 'unknown'], [1, second]]:
            self.assertIn(self.client.put(endpoint,json={'ids':invalid}).status_code, [400,409])
        self.post('/api/users',{'username':'staff','display_name':'工作人员','password':'format-staff-123','role':'staff'})
        self.post('/api/logout', {})
        self.assertEqual(self.client.put(endpoint,json={'ids':[second,first]}).status_code,401)
        self.post('/api/login',{'username':'staff','password':'format-staff-123'})
        self.assertEqual(self.client.put(endpoint,json={'ids':[second,first]}).status_code,403)
        self.post('/api/login',{'username':'admin','password':'format-password-123'})
        self.assertEqual(self.client.put(endpoint,json={'ids':[second,first]}).status_code,200)
        self.assertEqual([a['id'] for a in self.client.get(f'/api/events/{event}').json()['attachments']],[second,first])
        self.assertTrue(any(a['target']=='附件顺序' for a in self.client.get('/api/audit').json()))
        # Mimic a pre-upgrade database, retaining all business records.
        from backend.db import connect, init_db
        with connect() as db:
            db.execute('ALTER TABLE attachments DROP COLUMN sort_order')
        init_db(); init_db()
        with connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM attachments').fetchone()[0],2)
            self.assertIn('sort_order', {r['name'] for r in db.execute('PRAGMA table_info(attachments)')})
        self.assertEqual(self.client.get(f'/api/events/{event}').status_code,200)


if __name__ == '__main__':
    unittest.main()
