import io
import unittest
import zipfile
from datetime import date, timedelta
from pathlib import Path
from docx import Document
from PIL import Image
from support import isolated_app


class AcceptanceAudit(unittest.TestCase):
    def setUp(self):
        self.context = isolated_app()
        self.client, self.paths = self.context.__enter__()
        self.addCleanup(self.context.__exit__, None, None, None)
        self.post('/api/setup', {'username':'admin','display_name':'管理员','password':'audit-password-123'})
        self.post('/api/login', {'username':'admin','password':'audit-password-123'})
        self.building = self.post('/api/buildings', {'name':'18栋'})['id']
        self.unit = self.post('/api/units', {'building_id':self.building,'name':'1单元'})['id']
        self.elevator = self.post('/api/elevators', {'unit_id':self.unit,'name':'1号梯'})['id']
        self.today = date.today().isoformat()

    def post(self, path, data):
        r=self.client.post(path,json=data)
        self.assertEqual(r.status_code,200,r.text)
        return r.json()

    def event(self, **values):
        data={'elevator_id':self.elevator,'title':'例行检查','type':'维修','occurred_on':self.today,'completed':True}
        data.update(values)
        return self.post('/api/events',data)['id']

    def test_cost_mode_survives_edit(self):
        part = {'name':'门锁','quantity':1,'unit_price_yuan':'260'}
        event = self.event(parts=[part])
        record = self.client.get(f'/api/events/{event}').json()
        self.assertEqual((record['cost_auto'], record['cost_cents']), (1, 26000))
        changed = {**record, 'cost_yuan':'', 'parts':[{**part, 'unit_price_yuan':'300'}]}
        self.assertEqual(self.client.put(f'/api/events/{event}',json=changed).status_code,200)
        record = self.client.get(f'/api/events/{event}').json()
        self.assertEqual((record['cost_auto'], record['cost_cents']), (1, 30000))
        self.assertEqual(self.client.put(f'/api/events/{event}',json={**record, 'cost_yuan':'0', 'parts':[part]}).status_code,200)
        record = self.client.get(f'/api/events/{event}').json()
        self.assertEqual((record['cost_auto'], record['cost_cents']), (0, 0))

    def test_inline_elevator_location_and_edit(self):
        data = {'building_name':'20栋','unit_name':'2单元','name':'2号梯','equipment_number':'000021'}
        lift = self.post('/api/elevators',data)['id']
        self.assertEqual(self.client.get(f'/api/elevators/{lift}').json()['building_name'],'20栋')
        self.assertEqual(self.client.post('/api/elevators',json=data).status_code,409)
        self.assertEqual(self.client.put(f'/api/elevators/{lift}',json={**data,'building_name':'21栋'}).status_code,200)
        self.assertEqual(self.client.get(f'/api/elevators/{lift}').json()['building_name'],'21栋')
        self.assertEqual(self.client.post('/api/elevators',json={**data,'building_name':'','unit_name':'未建单元'}).status_code,400)

    def test_excel_preview_commit_repeat_and_export(self):
        from backend.archive_excel import workbook_bytes
        from openpyxl import load_workbook
        raw = workbook_bytes([{'building_name':'新楼栋','unit_name':'1单元','name':'1号梯','maintenance_phone':'00123','equipment_number':'000042','notes':'=SUM(1,2)'}])
        def upload(preview):
            return self.client.post('/api/elevator-archives/import',params={'dry_run':preview},files={'file':('电梯.xlsx',raw)}).json()
        self.assertEqual(upload(True)['added'],1)
        self.assertFalse(any(b['name']=='新楼栋' for b in self.client.get('/api/buildings').json()))
        self.assertEqual(upload(False)['added'],1)
        self.assertEqual(upload(False)['skipped'],1)
        records = self.client.get('/api/elevators').json()
        lift = next(e for e in records if e['building_name']=='新楼栋')
        exported = self.client.get('/api/elevator-archives/export',params={'building_id':lift['building_id']})
        workbook=load_workbook(io.BytesIO(exported.content))
        self.assertEqual(workbook['电梯档案']['E2'].value,'000042')
        self.assertEqual(workbook['电梯档案']['L2'].value,'00123')
        self.assertEqual(workbook['电梯档案']['P2'].data_type,'s')
        self.assertEqual(workbook['电梯档案'].max_row,2)
        workbook.close()

    def test_excel_invalid_batch_is_atomic(self):
        from backend.archive_excel import workbook_bytes
        base={'building_name':'不能留下的楼栋','unit_name':'1单元','name':'1号梯'}
        for records in ([base,{**base,'name':'2号梯','inspection_on':'2026-02-30'}],[base,base]):
            result=self.client.post('/api/elevator-archives/import?dry_run=false',files={'file':('批量.xlsx',workbook_bytes(records))})
            self.assertEqual(result.status_code,200,result.text)
            self.assertFalse(result.json()['valid'])
            self.assertFalse(any(b['name']=='不能留下的楼栋' for b in self.client.get('/api/buildings').json()))

    def test_excel_template_format_and_permissions(self):
        from openpyxl import load_workbook
        result=self.client.get('/api/elevator-archives/export?template=true')
        book=load_workbook(io.BytesIO(result.content))
        self.assertEqual(book['电梯档案']['L2'].number_format,'@')
        book['电梯档案']['A2']='18栋';book['电梯档案']['B2']='1单元';book['电梯档案']['C2']='=1+1'
        raw=io.BytesIO();book.save(raw);book.close()
        response=self.client.post('/api/elevator-archives/import',files={'file':('公式.xlsx',raw.getvalue())})
        self.assertEqual(response.status_code,400)
        self.assertIn('公式',response.json()['detail'])
        self.assertEqual(self.client.post('/api/elevator-archives/import',files={'file':('坏.xlsx',b'bad')}).status_code,400)
        self.post('/api/users',{'username':'import-staff','display_name':'工作人员','password':'test-password-123','role':'staff'})
        self.post('/api/logout',{})
        self.post('/api/login',{'username':'import-staff','password':'test-password-123'})
        self.assertEqual(self.client.post('/api/elevator-archives/import',files={'file':('表.xlsx',result.content)}).status_code,403)
        self.assertEqual(self.client.get('/api/elevator-archives/export').status_code,200)

    def test_webp_export_and_original_bytes(self):
        event=self.event()
        photo=io.BytesIO();Image.new('RGB',(500,250),'blue').save(photo,format='WEBP')
        r=self.client.post(f'/api/events/{event}/attachments',files={'file':('现场.webp',photo.getvalue(),'image/webp')})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(self.client.get('/api/attachments/'+r.json()['id']).content,photo.getvalue())
        r=self.client.get('/api/export',params={'scope':'elevator','elevator_id':self.elevator})
        self.assertEqual(r.status_code,200,r.text if r.status_code!=200 else '')
        self.assertEqual(len(Document(io.BytesIO(r.content)).inline_shapes),1)

    def test_state_follows_edit_and_not_backfill(self):
        recent=self.event(type='故障',title='当前故障',completed=False,recovered=False,elevator_status='故障')
        self.event(occurred_on='2020-01-01',title='旧维修')
        self.assertEqual(self.client.get(f'/api/elevators/{self.elevator}').json()['status'],'故障')
        r=self.client.put(f'/api/events/{recent}',json={'type':'故障','title':'当前故障','occurred_on':self.today,'completed':True,'recovered':True,'elevator_status':'正常'})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(self.client.get(f'/api/elevators/{self.elevator}').json()['status'],'正常')

    def test_search_all_keywords_same_record_and_company(self):
        self.event(title='门锁修复',company='精准维保有限公司')
        self.event(title='电机检查')
        self.assertFalse(self.client.get('/api/search',params={'q':'门锁 电机'}).json()['events'])
        results=self.client.get('/api/search',params={'q':'18栋 精准维保'}).json()['events']
        self.assertEqual(len(results),1)

    def test_validation_is_chinese_client_error(self):
        for value in ['Infinity','sNaN','NaN','-1','12.345']:
            r=self.client.post('/api/events',json={'elevator_id':self.elevator,'title':'非法费用','type':'维修','occurred_on':self.today,'cost_yuan':value})
            self.assertEqual(r.status_code,400,(value,r.text))
        r=self.client.post('/api/events',json={'elevator_id':self.elevator,'title':'日期越界','type':'配件更换','occurred_on':self.today,'parts':[{'name':'门锁','replaced_on':'9999-12-31','warranty_days':365}]})
        self.assertEqual(r.status_code,400,r.text)
        self.assertEqual(self.client.get('/api/events?elevator_id=abc').status_code,400)

    def test_deleted_elevator_records_hidden(self):
        event=self.event()
        self.assertEqual(self.client.delete(f'/api/elevators/{self.elevator}').status_code,200)
        self.assertEqual(self.client.get(f'/api/events/{event}').status_code,404)
        self.assertEqual(self.client.get('/api/dashboard').json()['stats']['monthly_repairs'],0)

    def test_export_contains_fault_times_and_completion(self):
        self.event(type='故障',title='3层故障',reported_at=self.today+'T09:12',arrived_at=self.today+'T09:30',restored_at=self.today+'T10:36',recovered=True,completed=True)
        r=self.client.get('/api/export',params={'scope':'elevator','elevator_id':self.elevator})
        self.assertEqual(r.status_code,200)
        text='\n'.join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
        for value in ('09:12','09:30','10:36','已完成','已恢复正常'):
            self.assertIn(value,text)

    def test_demo_can_be_recreated(self):
        from scripts.demo_data import install_demo,remove_demo
        install_demo();remove_demo();install_demo()
        self.assertTrue(any(x['name'].startswith('演示18栋') for x in self.client.get('/api/buildings').json()))

    def test_backup_can_restore_database_and_files(self):
        event=self.event()
        image=io.BytesIO();Image.new('RGB',(20,20),'red').save(image,format='PNG')
        attachment=self.client.post(f'/api/events/{event}/attachments',files={'file':('原图.png',image.getvalue(),'image/png')}).json()
        result=self.post('/api/backup',{})
        archive=self.client.get(result['download']).content
        import sqlite3
        with zipfile.ZipFile(io.BytesIO(archive)) as z:
            self.assertIsNone(z.testzip())
            restored=self.paths['DATA']/'restored.db';restored.write_bytes(z.read('data/elevator.db'))
            with sqlite3.connect(restored) as db:
                self.assertEqual(db.execute('PRAGMA integrity_check').fetchone()[0],'ok')
                row=db.execute('SELECT relative_path FROM attachments WHERE id=?',(attachment['id'],)).fetchone()
                self.assertEqual(z.read('storage/'+row[0]),image.getvalue())
                self.assertEqual(db.execute('SELECT COUNT(*) FROM events').fetchone()[0],1)
                self.assertEqual(db.execute('SELECT COUNT(*) FROM sessions').fetchone()[0],0)

    def test_edit_conflict_and_boolean_values(self):
        event=self.event(completed=True)
        first=self.client.get(f'/api/events/{event}').json()
        self.assertIs(first['completed'],True)
        changed={**first,'title':'新修改','cost_yuan':'0'}
        self.assertEqual(self.client.put(f'/api/events/{event}',json=changed).status_code,200)
        self.assertEqual(self.client.put(f'/api/events/{event}',json={**first,'title':'过期修改'}).status_code,409)
        self.assertEqual(self.client.get(f'/api/events/{event}').json()['title'],'新修改')

    def test_restore_and_attachment_permissions(self):
        from fastapi.testclient import TestClient
        from backend.main import app
        event=self.event()
        raw=io.BytesIO();Image.new('RGB',(20,10),'green').save(raw,format='PNG')
        photo=self.client.post(f'/api/events/{event}/attachments',files={'file':('现场.png',raw.getvalue(),'image/png')}).json()['id']
        staff=self.post('/api/users',{'username':'staff','display_name':'工作人员','role':'staff','password':'staff-password-123'})['id']
        other=TestClient(app)
        self.assertEqual(other.post('/api/login',json={'username':'staff','password':'staff-password-123'}).status_code,200)
        self.assertEqual(other.post(f'/api/events/{event}/attachments',files={'file':('现场.png',raw.getvalue(),'image/png')}).status_code,403)
        self.assertEqual(other.get('/api/trash').status_code,403)
        self.client.put(f'/api/users/{staff}',json={'active':False})
        self.client.put(f'/api/users/{staff}',json={'active':True})
        self.assertEqual(other.get('/api/me').status_code,401)
        self.assertEqual(self.client.delete(f'/api/events/{event}').status_code,200)
        self.assertEqual(self.client.get('/api/attachments/'+photo).status_code,404)
        self.assertEqual(self.client.post(f'/api/trash/events/{event}/restore').status_code,200)
        self.assertEqual(self.client.get('/api/attachments/'+photo).content,raw.getvalue())

    def test_missing_attachment_blocks_backup_and_export(self):
        event=self.event()
        raw=io.BytesIO();Image.new('RGB',(20,10),'green').save(raw,format='PNG')
        self.client.post(f'/api/events/{event}/attachments',files={'file':('现场.png',raw.getvalue(),'image/png')})
        next((self.paths['STORAGE']/'photos').glob('*')).unlink()
        self.assertEqual(self.client.post('/api/backup',json={}).status_code,400)
        self.assertEqual(self.client.get('/api/export',params={'scope':'all'}).status_code,400)
        self.assertFalse(list(self.paths['BACKUPS'].glob('*')))

    def test_photo_orientation_in_word(self):
        event=self.event()
        raw=io.BytesIO();photo=Image.new('RGB',(100,200),'blue');exif=photo.getexif();exif[274]=6;photo.save(raw,format='JPEG',exif=exif)
        item=self.client.post(f'/api/events/{event}/attachments',files={'file':('转向.jpg',raw.getvalue(),'image/jpeg')}).json()
        self.assertEqual(self.client.get('/api/attachments/'+item['id']).content,raw.getvalue())
        r=self.client.get('/api/export',params={'scope':'all'})
        image=Document(io.BytesIO(r.content)).inline_shapes[0]
        self.assertAlmostEqual(image.width/image.height,2.0,places=2)

    def test_date_and_type_filters_and_costs(self):
        self.event(type='日常维保',occurred_on='2026-01-01',cost_yuan='10')
        self.event(type='维修',occurred_on='2026-01-02',cost_yuan='20')
        self.event(type='维修',occurred_on='2026-01-03',cost_yuan='30')
        params={'start':'2026-01-01','end':'2026-01-02','type':'维修'}
        self.assertEqual(len(self.client.get('/api/events',params=params).json()),1)
        self.assertEqual(self.client.get('/api/costs',params=params).json()['cost_cents'],2000)
        for scope in ({'scope':'all'},{'scope':'unit','unit_id':self.unit}):
            r=self.client.get('/api/export',params={**scope,**params})
            document=Document(io.BytesIO(r.content))
            self.assertIn('1 条工作记录','\n'.join(p.text for p in document.paragraphs))
        self.assertEqual(self.client.get('/api/costs',params={'start':'2026-01-03','end':'2026-01-01'}).status_code,400)

    def test_request_validation_and_csrf(self):
        payload={'elevator_id':self.elevator,'title':'测试','type':'维修','occurred_on':self.today}
        for changes in ({'parts':[None]},{'type':[]},{'completed':'false'},{'cost_yuan':'Infinity'},{'title':{}},{'work_content':'bad\u0000text'},
                        {'started_at':self.today+'T10:00','ended_at':self.today+'T09:00'}):
            r=self.client.post('/api/events',json={**payload,**changes})
            self.assertEqual(r.status_code,400,r.text)
        r=self.client.post('/api/events',json=payload,headers={'origin':'http://evil.example'})
        self.assertEqual(r.status_code,403)
        self.assertEqual(self.client.get('/api/attachments/%2E%2E%2Frequirements.txt').status_code,404)
        self.assertEqual(self.client.get('/api/search',params={'q':"' OR 1=1 --"}).json()['events'],[])

    def test_demo_removal_protects_new_elevator(self):
        from scripts.demo_data import install_demo,remove_demo
        install_demo()
        units=self.client.get('/api/units').json()
        demo=next(x for x in units if x['building_name'].startswith('演示'))
        self.post('/api/elevators',{'unit_id':demo['id'],'name':'后来登记的电梯'})
        with self.assertRaises(SystemExit):
            remove_demo()
        self.assertTrue(any(x['name']=='后来登记的电梯' for x in self.client.get('/api/elevators').json()))

    def test_restart_from_restored_backup(self):
        import os, subprocess, sys, json
        event=self.event(title='重启后应保留的记录')
        backup=self.post('/api/backup',{})
        restored=self.paths['DATA'].parent/'restored'
        with zipfile.ZipFile(io.BytesIO(self.client.get(backup['download']).content)) as archive:
            archive.extractall(restored)
        environment=os.environ.copy()
        for key,folder in [('DATA','data'),('STORAGE','storage'),('EXPORT','exports'),('BACKUP','backups'),('LOG','logs')]:
            environment[f'ELEVATOR_{key}_DIR']=str(restored/folder)
        code="""from backend.main import app
from fastapi.testclient import TestClient
with TestClient(app) as c:
 r=c.post('/api/login',json={'username':'admin','password':'audit-password-123'})
 assert r.status_code==200,r.text
 records=c.get('/api/events').json()
 assert len(records)==1,records
 assert records[0]['title']=='重启后应保留的记录',records
 print('restart and restored login passed')
"""
        result=subprocess.run([sys.executable,'-c',code],env=environment,capture_output=True,text=True,timeout=20)
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)

if __name__=='__main__':
    unittest.main()
