"""Inspection archives: formal typography, traceable figures, original attachments."""
import csv
import io
import re
import uuid
import zipfile
from datetime import date
from decimal import Decimal

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageOps

from .db import COMMUNITY, EXPORTS, STORAGE, connect, rows

# Word/WPS use the reader's fonts, independent of the server operating system.
FONT = "仿宋"
TITLE_FONT = "方正小标宋简体"


def safe_name(value, max_length=100):
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", value).strip(" .")[:max_length]
    # Reserve space for the storage ID and extension on filesystems with 255-byte names.
    return name.encode("utf-8")[:220].decode("utf-8", errors="ignore").rstrip(" .")


def yuan(cents):
    return f"{Decimal(cents) / 100:.2f}元"


def friendly_date(value):
    if not value:
        return "未填写"
    parsed = date.fromisoformat(value)
    return f"{parsed.year}年{parsed.month}月{parsed.day}日"


def get_data(scope, building_id, unit_id, elevator_id, start, end, event_type):
    ids = {"building": building_id, "unit": unit_id, "elevator": elevator_id}
    if scope != "all" and not ids[scope]:
        raise ValueError("请先选择导出范围")
    sql = """SELECT e.*,u.name unit_name,u.building_id,b.name building_name FROM elevators e
      JOIN units u ON u.id=e.unit_id JOIN buildings b ON b.id=u.building_id
      WHERE e.deleted_at IS NULL AND u.deleted_at IS NULL AND b.deleted_at IS NULL"""
    args = []
    if scope == "building":
        sql += " AND b.id=?"; args.append(building_id)
    elif scope == "unit":
        sql += " AND u.id=?"; args.append(unit_id)
    elif scope == "elevator":
        sql += " AND e.id=?"; args.append(elevator_id)
    with connect() as db:
        db.execute("BEGIN")
        elevators = rows(db.execute(sql + " ORDER BY b.name,u.name,e.name,e.id", args))
        if not elevators:
            raise ValueError("范围内没有电梯")
        record_number = attachment_number = 0
        for elevator in elevators:
            query = "SELECT * FROM events WHERE elevator_id=? AND deleted_at IS NULL"
            params = [elevator["id"]]
            if start:
                query += " AND occurred_on>=?"; params.append(start)
            if end:
                query += " AND occurred_on<=?"; params.append(end)
            if event_type != "全部":
                query += " AND type=?"; params.append(event_type)
            # Backfilled events on the same day follow actual start time, not insertion ID.
            order = " ORDER BY occurred_on,COALESCE(NULLIF(fault_at,''),NULLIF(started_at,''),NULLIF(reported_at,''),'9999'),id"
            elevator["events"] = rows(db.execute(query + order, params))
            for event in elevator["events"]:
                record_number += 1
                event["record_number"] = record_number
                event["parts"] = rows(db.execute("SELECT * FROM event_parts WHERE event_id=? ORDER BY replaced_on,id", (event["id"],)))
                event["attachments"] = rows(db.execute("""SELECT * FROM attachments WHERE event_id=?
                    AND deleted_at IS NULL ORDER BY sort_order,created_at,id""", (event["id"],)))
                photo_number = 0
                for attachment in event["attachments"]:
                    if attachment["relative_path"].startswith("photos/"):
                        photo_number += 1
                        reference = f"图{record_number}-{photo_number}"
                    else:
                        attachment_number += 1
                        reference = f"附件{attachment_number}"
                    attachment["archive_reference"] = reference
    return elevators


def set_font(run, size=16, bold=False, face=FONT):
    run.font.name = face; run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0); run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + kind), face if kind == "eastAsia" else "Times New Roman")


def paragraph_format(p, line=28, indent=True, after=0):
    pf = p.paragraph_format
    pf.line_spacing = Pt(line); pf.space_before = Pt(0); pf.space_after = Pt(after)
    pf.first_line_indent = Pt(32 if indent else 0); pf.widow_control = True
    snap = OxmlElement("w:snapToGrid"); snap.set(qn("w:val"), "0")
    p._p.get_or_add_pPr().append(snap)


def add_text(document, text, size=16, bold=False, indent=True):
    p = document.add_paragraph(); paragraph_format(p, indent=indent)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_font(p.add_run(str(text)), size, bold)
    if bold:
        p.paragraph_format.keep_with_next = True
    return p


def add_heading(document, text, level=1):
    p = document.add_paragraph(style=f"Heading {level}")
    paragraph_format(p, indent=False, after=6)
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), 16, level == 3, "黑体" if level == 1 else "楷体" if level == 2 else FONT)
    return p


def cell_text(cell, value, header=False):
    cell.text = ""; p = cell.paragraphs[0]
    paragraph_format(p, line=20, indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    set_font(p.add_run(str(value if value not in (None, "") else "—")), 12, header, "黑体" if header else FONT)


def configure_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            item.set(qn("w:" + key), value)
        borders.append(item)
    table._tbl.tblPr.append(borders)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                item = OxmlElement("w:" + side)
                item.set(qn("w:w"), "100"); item.set(qn("w:type"), "dxa"); margins.append(item)
            cell._tc.get_or_add_tcPr().append(margins)


def add_table(document, headers, records, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, header in enumerate(headers):
        cell_text(table.rows[0].cells[i], header, True)
        shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "F2F2F2")
        table.rows[0].cells[i]._tc.get_or_add_tcPr().append(shade)
    for record in records:
        cells = table.add_row().cells
        for i, value in enumerate(record):
            cell_text(cells[i], value)
    configure_table(table, widths)
    paragraph_format(document.add_paragraph(), line=6, indent=False)
    return table


def add_photo(cell, image_path, caption):
    # Only the document copy is converted; the original and watermark remain untouched.
    drawing = io.BytesIO()
    with Image.open(image_path) as original:
        image = ImageOps.exif_transpose(original); image.thumbnail((2000, 2000))
        image.convert("RGB").save(drawing, format="PNG"); width, height = image.size
    drawing.seek(0); scale = min(Cm(7.0) / width, Cm(6.2) / height)
    p = cell.paragraphs[0]; paragraph_format(p, indent=False)
    # An exact line height would clip an inline drawing.
    p.paragraph_format.line_spacing = 1; p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(drawing, width=int(width * scale), height=int(height * scale))
    picture._inline.docPr.set("descr", caption.replace("\n", " "))
    p = cell.add_paragraph(style="Caption"); paragraph_format(p, line=18, indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(caption), 12, face="宋体")


def configure_document(document):
    section = document.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7); section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8); section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.5); section.footer_distance = Cm(2.0)
    for name, face, size in (("Normal", FONT, 16), ("Title", TITLE_FONT, 22),
                             ("Heading 1", "黑体", 16), ("Heading 2", "楷体", 16),
                             ("Heading 3", FONT, 16), ("Caption", "宋体", 12)):
        style = document.styles[name]
        style.font.name = face; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), face)
        pf = style.paragraph_format; pf.line_spacing = Pt(28); pf.space_before = Pt(0); pf.space_after = Pt(0)
        for border in style._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)


def cover_line(document, text, size=16, before=0, after=0, face=FONT, title=False):
    p = document.add_paragraph(style="Title" if title else None)
    paragraph_format(p, line=36 if title else 28, indent=False, after=after)
    p.paragraph_format.space_before = Pt(before); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), size, face=face)


def add_page_numbers(document, section):
    document.settings.odd_and_even_pages_header_footer = True
    section.footer.is_linked_to_previous = False; section.even_page_footer.is_linked_to_previous = False
    for footer, alignment in ((section.footer, WD_ALIGN_PARAGRAPH.RIGHT), (section.even_page_footer, WD_ALIGN_PARAGRAPH.LEFT)):
        p = footer.paragraphs[0]; paragraph_format(p, line=18, indent=False); p.alignment = alignment
        p.paragraph_format.left_indent = Pt(14); p.paragraph_format.right_indent = Pt(14)
        set_font(p.add_run("— "), 14, face="宋体")
        run = p.add_run(); set_font(run, 14, face="宋体")
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); run._r.addnext(field)
        set_font(p.add_run(" —"), 14, face="宋体")
    numbering = OxmlElement("w:pgNumType"); numbering.set(qn("w:start"), "1"); section._sectPr.append(numbering)


def build_document(elevators, start, end, photos, directory):
    document = Document(); configure_document(document)
    all_events = [event for e in elevators for event in e["events"]]
    buildings = list(dict.fromkeys(e["building_name"] for e in elevators))
    scope = (f"{elevators[0]['building_name']} {elevators[0]['unit_name']} {elevators[0]['name']}"
             if len(elevators) == 1 else "、".join(buildings))
    if len(scope) > 60:
        scope = f"{len(buildings)}栋楼  共{len(elevators)}部电梯"
    cover_line(document, COMMUNITY, 18, before=24, after=70, face="黑体")
    cover_line(document, "电梯工作记录档案", 22, after=28, face=TITLE_FONT, title=True)
    cover_line(document, scope, 16, after=36)
    effective_start = start or (min(e["occurred_on"] for e in all_events) if all_events else None)
    effective_end = end or (max(e["occurred_on"] for e in all_events) if all_events else None)
    cover_line(document, "统计周期")
    if effective_start and effective_end:
        period = f"{friendly_date(effective_start)} 至 {friendly_date(effective_end)}"
    elif effective_start:
        period = f"自{friendly_date(effective_start)}起"
    elif effective_end:
        period = f"截至{friendly_date(effective_end)}"
    else:
        period = "全部历史资料"
    cover_line(document, period, after=30)
    cover_line(document, f"共 {len(elevators)} 部电梯，{len(all_events)} 条工作记录。", 14, after=50)
    cover_line(document, f"编制单位：{COMMUNITY}", 14)
    cover_line(document, f"编制日期：{friendly_date(date.today().isoformat())}", 14)
    section = document.add_section(WD_SECTION_START.NEW_PAGE); add_page_numbers(document, section)
    add_heading(document, "编制说明")
    add_text(document, "本档案汇集所选范围和统计周期内的电梯基本资料、工作记录及相关附件，供工作检查、资料查阅和归档留存使用。")
    add_text(document, "按楼栋、单元、电梯分别整理。每部电梯内的工作按发生日期由早到晚排列；同日有时间的记录按故障发生、工作开始或报修时间排列，未填时间的记录按新增先后排列。")
    add_text(document, "工作记录在本册内连续编号。照片采用“图记录号-照片号”标注，如“图2-1”表示记录2的第1张照片。同一记录的照片按已保存的附件顺序排列，未调整时按上传先后排列。")
    add_text(document, "费用、质保期限及完成情况以登记内容为依据。未填写项目以“—”表示。原始附件在完整档案包中保存，图片仅为排版缩放，保留原始比例及水印。")
    for index, elevator in enumerate(elevators, 1):
        document.add_page_break()
        lift_label = f"{elevator['building_name']} {elevator['unit_name']} {elevator['name']}"
        add_heading(document, f"第{index}部电梯  {lift_label}"); add_heading(document, "（一）电梯基本情况", 2)
        pairs = [("所属小区", COMMUNITY), ("当前状态", elevator["status"]),
                 ("楼栋", elevator["building_name"]), ("单元", elevator["unit_name"]),
                 ("电梯", elevator["name"]), ("内部编号", elevator["internal_number"]),
                 ("设备编号", elevator["equipment_number"]), ("品牌", elevator["brand"]),
                 ("型号", elevator["model"]), ("出厂日期", elevator["manufactured_on"]),
                 ("投用日期", elevator["started_on"]), ("维保单位", elevator["maintenance_company"]),
                 ("维保联系人", elevator["maintenance_contact"]), ("联系电话", elevator["maintenance_phone"]),
                 ("下次维保", elevator["maintenance_due_on"]), ("下次年检", elevator["inspection_on"])]
        add_table(document, ["项目", "内容", "项目", "内容"],
                  [(pairs[i][0], pairs[i][1], pairs[i+1][0], pairs[i+1][1]) for i in range(0, len(pairs), 2)], [2.6, 5.2, 2.6, 5.2])
        if elevator["notes"]:
            add_text(document, "档案备注：" + elevator["notes"])
        summary_heading = add_heading(document, "（二）工作记录汇总", 2)
        if len(elevator["events"]) > 2:
            summary_heading.paragraph_format.page_break_before = True
        if not elevator["events"]:
            add_text(document, "所选时间段暂无工作记录。"); continue
        records = [(f"记录{e['record_number']}", e["occurred_on"], e["type"], e["title"],
                    "、".join(f"{p['name']}×{p['quantity']}" for p in e["parts"]) or "—", yuan(e["cost_cents"])) for e in elevator["events"]]
        add_table(document, ["编号", "日期", "类型", "工作内容", "更换配件", "费用"], records, [1.5, 2.4, 1.8, 4.5, 3.0, 2.4])
        add_text(document, f"本周期登记费用合计：{yuan(sum(e['cost_cents'] for e in elevator['events']))}")
        for event in elevator["events"]:
            document.add_page_break(); number = event["record_number"]
            add_heading(document, f"记录{number}  {event['type']}")
            for label, value in (("电梯", lift_label), ("发生日期", friendly_date(event["occurred_on"])),
                                 ("工作事项", event["title"]), ("完成情况", "已完成" if event["completed"] else "未完成")):
                add_text(document, f"{label}：{value}")
            if event["type"] == "故障":
                add_text(document, "发生停梯：" + ("是" if event["stopped"] else "否") + "；恢复情况：" + ("已恢复正常" if event["recovered"] else "尚未恢复"))
            for label, key in (("故障发生时间", "fault_at"), ("报修时间", "reported_at"), ("到场时间", "arrived_at"),
                               ("工作开始时间", "started_at"), ("工作结束时间", "ended_at"), ("恢复时间", "restored_at"),
                               ("工作内容", "work_content"), ("故障现象", "symptom"), ("故障原因", "cause"),
                               ("处理办法", "remedy"), ("处理情况", "resolution"), ("工作人员", "workers"),
                               ("维保或维修单位", "company"), ("备注", "notes")):
                if event[key]:
                    value = event[key].replace("T", " ") if key.endswith("_at") else event[key]
                    add_text(document, f"{label}：{value}")
            if event["parts"]:
                add_heading(document, "更换配件及质保情况", 2)
                for part in event["parts"]:
                    add_text(document, f"配件：{part['name']}；型号：{part['model'] or '—'}；数量：{part['quantity']}")
                    add_table(document, ["项目", "内容", "项目", "内容"], [
                        ("更换日期", part['replaced_on'], "单价", yuan(part['unit_price_cents'])),
                        ("配件合计", yuan(part['quantity'] * part['unit_price_cents']), "质保天数", f"{part['warranty_days']}天"),
                        ("质保开始", part['warranty_from'] or part['replaced_on'], "质保截止", part['warranty_until'] or '未设置')
                    ], [2.6, 5.2, 2.6, 5.2])
                    for key, label in (("reason", "更换原因"), ("brand", "厂家或品牌"), ("notes", "配件备注")):
                        if part[key]:
                            add_text(document, f"{label}：{part[key]}")
            add_text(document, f"记录费用：{yuan(event['cost_cents'])}")
            pictures = [a for a in event["attachments"] if a["relative_path"].startswith("photos/")]
            if photos:
                for offset in range(0, len(pictures), 4):
                    document.add_page_break()
                    add_heading(document, f"记录{number}  现场照片" + (f"  续{offset//4}" if offset else ""))
                    add_text(document, f"{lift_label}  {event['occurred_on']}", size=14, indent=False)
                    add_text(document, event["title"], size=14, indent=False)
                    batch = pictures[offset:offset+4]
                    table = document.add_table(rows=(len(batch)+1)//2, cols=2); configure_table(table, [7.8, 7.8])
                    for position, attachment in enumerate(batch):
                        # Opaque original names (including WeChat URLs) never appear in captions.
                        caption = f"{attachment['archive_reference']}  {attachment['category']}\n{event['occurred_on']}"
                        add_photo(table.rows[position//2].cells[position%2], STORAGE / attachment["relative_path"], caption)
        files = [(a["archive_reference"], f"记录{e['record_number']}", e["occurred_on"], a["category"], a["original_name"])
                 for e in elevator["events"] for a in e["attachments"] if not a["relative_path"].startswith("photos/")]
        if directory and files:
            document.add_page_break(); add_heading(document, "附件目录"); add_text(document, "电梯：" + lift_label)
            add_table(document, ["附件号", "记录号", "日期", "资料类别", "文件名称"], files, [1.7, 1.9, 2.4, 2.4, 7.2])
    properties = document.core_properties
    properties.title = "电梯工作记录档案"; properties.author = COMMUNITY
    properties.last_modified_by = COMMUNITY; properties.comments = ""
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); document.settings.element.append(update)
    return document


def make_export(scope, building_id, unit_id, elevator_id, start, end, event_type, photos, directory, package):
    elevators = get_data(scope, building_id, unit_id, elevator_id, start, end, event_type)
    label = "全部电梯" if scope == "all" else (
        elevators[0]["building_name"] if scope == "building" else
        elevators[0]["building_name"] + elevators[0]["unit_name"] if scope == "unit" else
        elevators[0]["building_name"] + elevators[0]["unit_name"] + elevators[0]["name"])
    date_label = f"{(start or '不限').replace('-', '')}-{(end or '不限').replace('-', '')}"
    name = safe_name(f"{COMMUNITY}_{label}_{date_label}_电梯工作档案", 180)
    unique = uuid.uuid4().hex[:12]; docx_path = EXPORTS / f"{unique}_{name}.docx"
    for elevator in elevators:
        for event in elevator["events"]:
            for attachment in event["attachments"]:
                path = (STORAGE / attachment["relative_path"]).resolve()
                if not path.is_relative_to(STORAGE.resolve()) or not path.is_file():
                    raise ValueError(f"附件缺失，无法生成完整档案：{attachment['original_name']}")
    build_document(elevators, start, end, photos, directory).save(docx_path)
    if not package:
        return docx_path, f"{name}.docx"
    zip_path = EXPORTS / f"{unique}_{name}.zip"
    manifest = io.StringIO(newline=""); writer = csv.writer(manifest)
    writer.writerow(["资料编号", "电梯", "记录号", "事件日期", "资料类别", "原始文件名", "包内路径"])
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(docx_path, f"{name}.docx")
        for elevator in elevators:
            for event in elevator["events"]:
                for attachment in event["attachments"]:
                    path = STORAGE / attachment["relative_path"]
                    folder = safe_name(f"{elevator['building_name']}{elevator['unit_name']}{elevator['name']}", 60)
                    filename = safe_name(attachment["original_name"], 70); reference = attachment["archive_reference"]
                    if attachment["relative_path"].startswith("photos/"):
                        filename = safe_name(attachment["category"], 30) + path.suffix
                    archive_name = f"附件/{folder}/{reference}_{filename}"
                    archive.write(path, archive_name)
                    values = [reference, folder, f"记录{event['record_number']}", event["occurred_on"], attachment["category"], attachment["original_name"], archive_name]
                    # Untrusted names must not become executable spreadsheet formulas.
                    writer.writerow(["'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value for value in values])
        archive.writestr("原始附件清单.csv", manifest.getvalue().encode("utf-8-sig"))
    return zip_path, f"{name}.zip"
