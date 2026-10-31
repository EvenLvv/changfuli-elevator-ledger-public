"""Editable archive snapshots. Business records and original files remain read-only."""
import csv
import io
import json
import uuid
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from PIL import Image, ImageOps

from . import db as store
from .exporter import (add_heading, add_page_numbers, add_table, add_text,
                       configure_document, cover_line, paragraph_format,
                       safe_name, set_font, yuan)


def source_attachments(source):
    return [(elevator, event, attachment) for elevator in source for event in elevator['events']
            for attachment in event['attachments']]


def initial_blocks(source, photos=True):
    blocks = []

    def add(kind, text=None, **values):
        blocks.append({'id': uuid.uuid4().hex, 'type': kind, **({'text': str(text)} if text is not None else {}), **values})

    def fields(item, pairs):
        return '；'.join(f'{name}：{item[key]}' for key, name in pairs if item.get(key) not in (None, ''))

    for elevator in source:
        label = f"{elevator['building_name']} {elevator['unit_name']} {elevator['name']}"
        add('heading', label)
        for pairs in ((('internal_number', '内部编号'), ('equipment_number', '设备编号'), ('brand', '品牌'),
                       ('model', '型号'), ('manufactured_on', '出厂日期'), ('started_on', '投用日期')),
                      (('maintenance_company', '维保单位'), ('maintenance_contact', '维保联系人'),
                       ('maintenance_phone', '联系电话'), ('inspection_on', '下次年检'), ('maintenance_due_on', '下次维保')),
                      (('status', '当前状态'), ('notes', '电梯档案备注'))):
            text = fields(elevator, pairs)
            if text:
                add('paragraph', text)
        if not elevator['events']:
            add('paragraph', '所选时间段暂无工作记录。')
        for event in elevator['events']:
            add('heading', f"{event['occurred_on']}  {event['type']}  {event['title']}")
            add('paragraph', f"电梯：{label}\n发生日期：{event['occurred_on']}\n工作事项：{event['title']}\n完成情况：{'已完成' if event['completed'] else '未完成'}")
            if event['type'] == '故障':
                add('paragraph', f"发生停梯：{'是' if event['stopped'] else '否'}；恢复情况：{'已恢复正常' if event['recovered'] else '尚未恢复'}")
            times = fields({key: value.replace('T', ' ') if isinstance(value, str) else value for key, value in event.items()},
                           (('fault_at', '故障发生时间'), ('reported_at', '报修时间'), ('arrived_at', '到场时间'),
                            ('started_at', '工作开始时间'), ('ended_at', '工作结束时间'), ('restored_at', '恢复时间')))
            if times:
                add('paragraph', times)
            for key, name in (('work_content', '工作内容'), ('symptom', '故障现象'), ('cause', '故障原因'),
                              ('remedy', '处理办法'), ('resolution', '处理情况'), ('notes', '工作备注')):
                if event.get(key):
                    add('paragraph', f'{name}：{event[key]}')
            people = fields(event, (('workers', '工作人员'), ('company', '维保或维修单位')))
            if people:
                add('paragraph', people)
            for part in event['parts']:
                add('heading', f"更换配件：{part['name']}")
                details = fields(part, (('model', '型号'), ('quantity', '数量'), ('replaced_on', '更换日期'), ('brand', '厂家或品牌')))
                add('paragraph', details + f"；配件单价：{yuan(part['unit_price_cents'])}；配件合计：{yuan(part['unit_price_cents'] * part['quantity'])}")
                warranty = {**part, 'warranty_from': part.get('warranty_from') or part['replaced_on']}
                add('paragraph', fields(warranty, (('reason', '更换原因'), ('warranty_days', '质保天数'),
                                       ('warranty_from', '质保开始日期'), ('warranty_until', '质保截止日期'), ('notes', '配件备注'))))
            add('paragraph', f"记录费用：{yuan(event['cost_cents'])}")
            if photos:
                for attachment in event['attachments']:
                    if attachment['relative_path'].startswith('photos/'):
                        add('image', attachment_id=attachment['id'], caption=attachment['category'])
    return blocks


def draft_metadata(source, filters):
    first = source[0]
    buildings = list(dict.fromkeys(e['building_name'] for e in source))
    if len(source) == 1:
        scope = f"{first['building_name']} {first['unit_name']} {first['name']}"
    else:
        scope = {'building': first['building_name'], 'unit': f"{first['building_name']} {first['unit_name']}"}.get(filters['scope'], '全部电梯')
        scope += f"（{len(source)}部电梯）"
    if len(scope) > 60:
        scope = f'{len(buildings)}栋楼 共{len(source)}部电梯'
    start, end = filters['start'], filters['end']
    period = f'{start} 至 {end}' if start and end else f'自{start}起' if start else f'截至{end}' if end else '全部历史资料'
    return {'community': store.COMMUNITY, 'scope_label': scope, 'period_label': period}


def source_status(source):
    """Check the whole soft-delete chain, using the server snapshot's immutable file references."""
    issues, status = [], {}
    with store.connect() as db:
        db.execute('BEGIN')
        for elevator in source:
            live_elevator = db.execute('''SELECT e.id FROM elevators e JOIN units u ON u.id=e.unit_id
                JOIN buildings b ON b.id=u.building_id WHERE e.id=? AND e.deleted_at IS NULL
                AND u.deleted_at IS NULL AND b.deleted_at IS NULL''', (elevator['id'],)).fetchone()
            if not live_elevator:
                issues.append(f"来源电梯或所属楼栋、单元已删除：{elevator['building_name']} {elevator['unit_name']} {elevator['name']}")
            for event in elevator['events']:
                live = db.execute('''SELECT v.id FROM events v JOIN elevators e ON e.id=v.elevator_id
                    JOIN units u ON u.id=e.unit_id JOIN buildings b ON b.id=u.building_id
                    WHERE v.id=? AND v.elevator_id=? AND v.deleted_at IS NULL AND e.deleted_at IS NULL
                    AND u.deleted_at IS NULL AND b.deleted_at IS NULL''', (event['id'], elevator['id'])).fetchone()
                if not live:
                    issues.append(f"来源工作记录已删除或不可用：{event['title']}")
                for attachment in event['attachments']:
                    reason = ''
                    current = db.execute('''SELECT relative_path FROM attachments
                        WHERE id=? AND event_id=? AND elevator_id=? AND deleted_at IS NULL''',
                        (attachment['id'], event['id'], elevator['id'])).fetchone()
                    path = (store.STORAGE / attachment['relative_path']).resolve()
                    if not live:
                        reason = '来源工作记录或所属楼栋、单元、电梯已删除'
                    elif not current or current['relative_path'] != attachment['relative_path']:
                        reason = '来源附件已删除或已变化'
                    elif not path.is_relative_to(store.STORAGE.resolve()) or not path.is_file():
                        reason = '来源附件文件缺失或路径无效'
                    status[attachment['id']] = reason
                    if reason:
                        issues.append(f"{reason}：{attachment['original_name']}")
    return status, issues


def attachment_directory(source, blocks, status):
    references = {block['attachment_id']: f'图{number}' for number, block in
                  enumerate((b for b in blocks if b['type'] == 'image'), 1)}
    entries, file_number = [], 0
    for elevator, event, attachment in source_attachments(source):
        photo = attachment['relative_path'].startswith('photos/')
        if not photo:
            file_number += 1
        reason = status.get(attachment['id'], '')
        entries.append({'id': attachment['id'], 'event_id': event['id'], 'event_title': event['title'],
                        'event_date': event['occurred_on'],
                        'elevator_label': f"{elevator['building_name']} {elevator['unit_name']} {elevator['name']}",
                        'category': attachment['category'], 'original_name': attachment['original_name'], 'source_order': len(entries),
                        'display_name': f"{event['occurred_on']} {attachment['category']}" if photo else attachment['original_name'],
                        'reference': references.get(attachment['id'], '未编入正文' if photo else f'附件{file_number}'),
                        'is_photo': photo, 'in_body': attachment['id'] in references,
                        'available': not reason, 'unavailable_reason': reason})
    # The article's picture order is also the directory and package order.
    order = {block['attachment_id']: number for number, block in enumerate(b for b in blocks if b['type'] == 'image')}
    return sorted(entries, key=lambda entry: order.get(entry['id'], len(order)))


def draft_detail(row):
    source, blocks = json.loads(row['source_json']), json.loads(row['blocks_json'])
    filters = json.loads(row['filters_json'])
    status, issues = source_status(source)
    directory = attachment_directory(source, blocks, status)
    images = [{**entry, 'url': f"/api/export-drafts/{row['id']}/images/{entry['id']}"}
              for entry in directory if entry['is_photo']]
    return {key: row[key] for key in ('id', 'created_by', 'created_by_name', 'title', 'revision', 'created_at', 'updated_at')} | {
        'filters': filters, 'metadata': draft_metadata(source, filters), 'blocks': blocks, 'images': images,
        'attachment_directory': directory, 'source_warnings': issues}


def build_draft_document(title, blocks, source, directory, metadata):
    document = Document(); configure_document(document)
    cover_line(document, title, 22, after=18, face='方正小标宋简体', title=True)
    add_text(document, '编制单位：' + metadata['community'], size=12, indent=False)
    add_text(document, '资料范围：' + metadata['scope_label'], size=12, indent=False)
    add_text(document, '统计周期：' + metadata['period_label'], size=12, indent=False)
    add_page_numbers(document, document.sections[0])
    attachments = {attachment['id']: attachment for _, _, attachment in source_attachments(source)}
    image_number = 0
    for block in blocks:
        if block['type'] == 'heading':
            add_heading(document, block['text'])
        elif block['type'] == 'paragraph':
            paragraph = add_text(document, block['text'])
            if '\n' in block['text'] or '\r' in block['text']:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            image_number += 1
            drawing = io.BytesIO()
            with Image.open(store.STORAGE / attachments[block['attachment_id']]['relative_path']) as original:
                image = ImageOps.exif_transpose(original); image.thumbnail((2000, 2000))
                image.convert('RGB').save(drawing, format='PNG'); width, height = image.size
            drawing.seek(0); scale = min(Cm(14.8) / width, Cm(12) / height)
            paragraph = document.add_paragraph(); paragraph_format(paragraph, indent=False)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.keep_with_next = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = f"图{image_number}" + (f"  {block['caption']}" if block['caption'] else '')
            picture = paragraph.add_run().add_picture(drawing, width=int(width * scale), height=int(height * scale))
            picture._inline.docPr.set('descr', caption)
            paragraph = document.add_paragraph(style='Caption'); paragraph_format(paragraph, line=18, indent=False, after=8)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_font(paragraph.add_run(caption), 12, face='宋体')
    if directory:
        add_heading(document, '附件目录')
        add_table(document, ['资料编号', '来源记录', '日期', '资料类别', '文件名称'],
                  [(entry['reference'], entry['elevator_label'] + '\n' + entry['event_title'], entry['event_date'], entry['category'], entry['display_name'])
                   for entry in directory], [1.6, 4.0, 2.4, 2.4, 5.2])
    document.core_properties.title = title; document.core_properties.author = store.COMMUNITY
    document.core_properties.last_modified_by = store.COMMUNITY; document.core_properties.comments = ''
    return document


def make_draft_export(row, package, include_directory):
    source, blocks = json.loads(row['source_json']), json.loads(row['blocks_json'])
    status, issues = source_status(source)
    if issues:
        raise ValueError(issues[0] + '；请恢复来源资料，或重新生成文稿')
    directory = attachment_directory(source, blocks, status)
    name = safe_name(row['title'], 160) or '工作文稿'
    prefix = f"draft-{row['id']}-{uuid.uuid4().hex[:12]}_"
    docx_path = store.EXPORTS / f'{prefix}{name}.docx'
    zip_path = store.EXPORTS / f'{prefix}{name}.zip'
    temporary = docx_path.with_suffix('.partial')
    zip_temporary = zip_path.with_suffix('.partial.zip')
    try:
        build_draft_document(row['title'], blocks, source, directory if include_directory else [],
                             draft_metadata(source, json.loads(row['filters_json']))).save(temporary)
        temporary.replace(docx_path)
        if package:
            original = {attachment['id']: attachment for _, _, attachment in source_attachments(source)}
            manifest = io.StringIO(newline=''); writer = csv.writer(manifest)
            writer.writerow(['资料编号', '电梯', '来源记录', '事件日期', '资料类别', '原始文件名', '正文状态', '包内路径'])
            with zipfile.ZipFile(zip_temporary, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
                archive.write(docx_path, name + '.docx')
                for number, entry in enumerate(directory, 1):
                    attachment = original[entry['id']]; path = store.STORAGE / attachment['relative_path']
                    filename = safe_name(entry['category'], 30) + path.suffix if entry['is_photo'] else safe_name(entry['original_name'], 70)
                    archive_name = f"附件/{number:04d}_{entry['reference']}_{filename}"
                    archive.write(path, archive_name)
                    values = [entry['reference'], entry['elevator_label'], entry['event_title'], entry['event_date'],
                              entry['category'], entry['original_name'], '已编入正文' if entry['in_body'] else '未编入正文', archive_name]
                    writer.writerow(["'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value for value in values])
                archive.writestr('原始附件清单.csv', manifest.getvalue().encode('utf-8-sig'))
            zip_temporary.replace(zip_path)
        return docx_path, zip_path if package else None, name + ('.zip' if package else '.docx')
    except Exception:
        docx_path.unlink(missing_ok=True); zip_path.unlink(missing_ok=True)
        raise
    finally:
        temporary.unlink(missing_ok=True); zip_temporary.unlink(missing_ok=True)
