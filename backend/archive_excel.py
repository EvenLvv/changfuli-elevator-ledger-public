"""Bounded, round-trippable base-archive spreadsheets (no work records)."""
import io
from datetime import date, datetime
from zipfile import ZipFile, BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from .db import COMMUNITY

COLUMNS = [('building_name','楼栋'),('unit_name','单元'),('name','电梯名称'),
           ('internal_number','内部编号'),('equipment_number','设备编号'),('brand','品牌'),('model','型号'),
           ('manufactured_on','出厂日期'),('started_on','投用日期'),('maintenance_company','维保单位'),
           ('maintenance_contact','维保联系人'),('maintenance_phone','联系电话'),
           ('maintenance_due_on','下次维保日期'),('inspection_on','下次年检日期'),('status','状态'),('notes','备注')]


def workbook_bytes(records):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '电梯档案'
    sheet.append([label for _, label in COLUMNS])
    for record in records:
        sheet.append([record.get(key) or '' for key, _ in COLUMNS])
        for cell in sheet[sheet.max_row]:
            # Explicit text preserves leading zeroes and prevents Excel formula execution.
            cell.data_type = 's'
            cell.number_format = '@'
    for cell in sheet[1]:
        cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='234D70')
    for index, (key, _) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 30 if key in ('notes','maintenance_company') else 20
        if not records:
            for row in range(2, 502):
                sheet.cell(row, index).number_format = '@'
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    validation = DataValidation(type='list', formula1='"正常,故障,维修中,停梯,待处理"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add('O2:O501')
    instructions = workbook.create_sheet('填写说明')
    for line in [COMMUNITY + ' · 电梯基础档案', '在“电梯档案”页每行填写一部电梯，楼栋、单元、电梯名称为必填。',
                 '例如：18栋 / 1单元 / 1号梯。名称相同才会归入同一楼栋和单元，请保持一致。',
                 '日期格式：2026-09-27；设备编号和电话按文本填写以保留开头的0。',
                 '每次最多500部，文件不超过5MB，仅支持.xlsx，不支持公式或宏。',
                 '导入只新增档案；相同楼栋、单元、电梯名称的已有档案会跳过，不覆盖。',
                 '请勿改动首行列名。工作记录、照片和附件不包含在本表中。']:
        instructions.append([line])
    instructions.column_dimensions['A'].width = 100
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def read_records(raw):
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024 or len(archive.infolist()) > 200:
                raise ValueError('表格内容过大，请拆分为每次最多500部电梯')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError('无法读取表格，请使用下载的.xlsx模板') from exc
    try:
        if '电梯档案' not in workbook.sheetnames:
            raise ValueError('缺少“电梯档案”工作表，请使用下载的模板')
        sheet = workbook['电梯档案']
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(max_col=len(COLUMNS)+1)
        header = next(iterator, ())
        if [c.value for c in header[:len(COLUMNS)]] != [label for _, label in COLUMNS]:
            raise ValueError('表头不匹配，请下载模板并保留首行列名')
        result = []
        for number, row in enumerate(iterator, 2):
            if number > 501:
                raise ValueError('每次最多500行，请拆分表格后导入')
            if not any(cell.value is not None for cell in row):
                continue
            if row[-1].value is not None:
                raise ValueError(f'第{number}行有多余列，请保留模板结构')
            record = {}
            for (key, label), cell in zip(COLUMNS, row):
                if cell.data_type in ('f', 'e'):
                    raise ValueError(f'第{number}行“{label}”不能使用公式或错误值')
                value = cell.value
                if isinstance(value, (date, datetime)):
                    value = value.strftime('%Y-%m-%d')
                record[key] = str(value).strip() if value is not None else ''
            result.append((number, record))
        if not result:
            raise ValueError('表格没有电梯档案，请从第2行开始填写')
        return result
    finally:
        workbook.close()
